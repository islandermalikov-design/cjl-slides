"""Правка 1: руководитель (Поддубная) — в рамку на титуле. Правка 2: подписи категорий на слайде корреляций — отдельными текстовыми блоками,
чтобы PowerPoint не заменял длинные подписи многоточием."""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
EMU = 914400.0
I = lambda v: Emu(int(round(v * EMU)))
TXT = RGBColor(0x2F, 0x3E, 0x55)

prs = Presentation(sys.argv[1])

# ---- 1. титульный слайд
s = prs.slides[0]
box = [sh for sh in s.shapes if sh.name == "Соискатель и руководители"][0]
paras = box.text_frame.paragraphs
idx = [i for i, p in enumerate(paras) if "Поддубная" in p.text]
assert idx, "строка с Поддубной не найдена"
p_el = paras[idx[0]]._p
src_runs = paras[idx[0]].runs
size = src_runs[0].font.size
text = paras[idx[0]].text
# убираем строку из общего блока
p_el.getparent().remove(p_el)
# отдельный блок в рамке
y = box.top / EMU + 0.30 * idx[0] + 0.03
tb = s.shapes.add_textbox(I(box.left / EMU - 0.08), I(y), I(6.4), I(0.42))
tb.name = "Руководитель в рамке"
tf = tb.text_frame
tf.word_wrap = False
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = tf.margin_right = I(0.08)
tf.margin_top = tf.margin_bottom = I(0.02)
r = tf.paragraphs[0].add_run()
r.text = text
r.font.size = size or Pt(18)
r.font.italic = True
r.font.color.rgb = TXT
tb.line.color.rgb = TXT
tb.line.width = Pt(1.5)

# ---- 2. слайд корреляций: подписи слева
for s in prs.slides:
    if not any(sh.name == "Корреляции Спирмена" for sh in s.shapes):
        continue
    ch_sh = [sh for sh in s.shapes if sh.name == "Корреляции Спирмена"][0]
    cs = ch_sh.chart._chartSpace
    cats = list(ch_sh.chart.plots[0].categories)[::-1]  # сверху вниз
    lay = cs.find(".//c:plotArea/c:layout/c:manualLayout", NS)
    g = lambda t: float(lay.find("c:" + t, NS).get("val"))
    fx, fy, fw, fh = ch_sh.left / EMU, ch_sh.top / EMU, ch_sh.width / EMU, ch_sh.height / EMU
    px, py, pw, ph = fx + g("x") * fw, fy + g("y") * fh, g("w") * fw, g("h") * fh
    # скрыть штатные подписи категорий
    cat_ax = cs.find(".//c:plotArea/c:catAx", NS)
    cat_ax.find("c:tickLblPos", NS).set("val", "none")
    # заголовок оси категорий — отдельным вертикальным блоком слева
    ax_title = cat_ax.find("c:title", NS)
    if ax_title is not None:
        cat_ax.remove(ax_title)
    atb = s.shapes.add_textbox(I(fx + 0.35 - 1.9), I(py + ph / 2 - 0.2), I(3.8), I(0.4))
    atb.name = "Заголовок оси категорий"
    atb.rotation = 270
    atf = atb.text_frame
    atf.word_wrap = True
    atf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ap = atf.paragraphs[0]
    ap.alignment = PP_ALIGN.CENTER
    ar = ap.add_run(); ar.text = "Пары показателей"; ar.font.size = Pt(14); ar.font.color.rgb = TXT
    slot = ph / len(cats)
    left = fx + 0.6
    for i, label in enumerate(cats):
        t = s.shapes.add_textbox(I(left), I(py + i * slot), I(px - left - 0.08), I(slot))
        t.name = "Подпись категории %d" % (i + 1)
        tf = t.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.RIGHT
        run = para.add_run()
        run.text = label
        run.font.size = Pt(16)
        run.font.color.rgb = TXT
prs.save(sys.argv[2])
