"""Правки к версии 17: подписи категорий на слайдах 33-34 — отдельными блоками; блоки на слайдах 2 и 6 — под ссылки."""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
EMU = 914400.0
I = lambda v: Emu(int(round(v * EMU)))
TXT = RGBColor(0x2F, 0x3E, 0x55)
prs = Presentation(sys.argv[1])


def place(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)


def by_name(s, n):
    return [sh for sh in s.shapes if sh.name == n]


def find(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix) and sh.top < I(1.2):
                return s


# ---------- слайды 33-34: подписи категорий вне диаграммы ----------
def outside_labels(slide, chart_shape, axis_title):
    cs = chart_shape.chart._chartSpace
    cats = list(chart_shape.chart.plots[0].categories)[::-1]  # сверху вниз
    lay = cs.find(".//c:plotArea/c:layout/c:manualLayout", NS)
    g = lambda t: float(lay.find("c:" + t, NS).get("val"))
    fx, fy, fw, fh = chart_shape.left / EMU, chart_shape.top / EMU, chart_shape.width / EMU, chart_shape.height / EMU
    px, py, pw, ph = fx + g("x") * fw, fy + g("y") * fh, g("w") * fw, g("h") * fh
    cat_ax = cs.find(".//c:plotArea/c:catAx", NS)
    cat_ax.find("c:tickLblPos", NS).set("val", "none")
    t = cat_ax.find("c:title", NS)
    if t is not None:
        cat_ax.remove(t)
    # заголовок оси — вертикальный блок слева
    atb = slide.shapes.add_textbox(I(fx + 0.3 - 1.9), I(py + ph / 2 - 0.2), I(3.8), I(0.4))
    atb.name = "Заголовок оси категорий"
    atb.rotation = 270
    atb.text_frame.word_wrap = True
    atb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    ap = atb.text_frame.paragraphs[0]; ap.alignment = PP_ALIGN.CENTER
    ar = ap.add_run(); ar.text = axis_title; ar.font.size = Pt(14); ar.font.color.rgb = TXT
    slot = ph / len(cats)
    left = fx + 0.55
    for i, label in enumerate(cats):
        tb = slide.shapes.add_textbox(I(left), I(py + i * slot), I(px - left - 0.08), I(slot))
        tb.name = "Подпись категории %d" % (i + 1)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]; para.alignment = PP_ALIGN.RIGHT
        run = para.add_run(); run.text = label; run.font.size = Pt(16); run.font.color.rgb = TXT


for prefix, title in (("Факторы, определяющие", "Фактор"), ("Сопряженность", "Исход лечения")):
    s = find(prefix)
    if s is None:
        continue
    for sh in [x for x in s.shapes if x.has_chart]:
        outside_labels(s, sh, title)

# ---------- слайд 2: блоки под ссылки ----------
s = find("Актуальность")
for n in ("Показатель 1", "Показатель 2"):
    for sh in by_name(s, n): place(sh, y=1.45, h=1.25)
for n in ("Факт 1", "Факт 2", "Факт 3"):
    for sh in by_name(s, n): place(sh, y=2.8, h=1.3)
for sh in by_name(s, "Вывод по актуальности"): place(sh, y=4.2, h=1.0)
for n in ("Пустой блок 1", "Пустой блок 2"):
    for sh in by_name(s, n): place(sh, y=5.3, h=2.05)

# ---------- слайд 6: колонки и блоки ----------
s = find("Дизайн исследования")
cols = [(0.5, 4.7), (5.44, 3.55), (9.23, 3.6)]
for sh in by_name(s, "База 1") + by_name(s, "База 2"): place(sh, y=1.05, h=0.95)
for sh in by_name(s, "Рандомизация"): place(sh, y=2.08, h=0.42)
for k, (x, w) in enumerate(cols, 1):
    grp = [sh for sh in s.shapes if sh.name in ("Группа контроля (n=43)", "Группа сравнения (n=47)", "Основная группа (n=45)")][k - 1]
    place(grp, x=x, y=2.62, w=w, h=0.36)
    subs = [sh for sh in s.shapes if sh.name.startswith("Подгруппа %d" % k)]
    subs.sort(key=lambda z: z.name)
    for j, sub in enumerate(subs):
        place(sub, x=x + j * (w / 2), y=2.98, w=w / 2 - 0.02, h=0.6)
    for sh in by_name(s, "Метод %d" % k): place(sh, x=x, y=3.8, w=w, h=2.2)
# стрелки
centers = [x + w / 2 for x, w in cols]
for sh in by_name(s, "Стрелка"):
    y = sh.top / EMU; x = sh.left / EMU
    if abs(y - 2.15) < 0.03:
        place(sh, y=2.0, h=0.08)
    elif abs(y - 2.72) < 0.03 or abs(y - 3.87) < 0.03:
        k = min(range(3), key=lambda i: abs(centers[i] - x if False else (0.5 + [0, 4.19, 8.38][i] + 1.975) - x))
        if abs(y - 2.72) < 0.03: place(sh, x=centers[k], y=2.5, h=0.12)
        else: place(sh, x=centers[k], y=3.58, h=0.22)
# шрифты ссылок: имена 14, длинное описание источника 11
for n in ("Метод 1", "Метод 2", "Метод 3"):
    for sh in by_name(s, n):
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.font.name == "Times New Roman":
                    sz = r.font.size.pt if r.font.size else 16
                    r.font.size = Pt(11 if sz <= 10 else 14)
# слайд 34: заголовок в 3 строки налезает на «Положение 3» -> 20 pt
s = find("Сопряженность")
if s is not None:
    for sh in s.shapes:
        if sh.is_placeholder and sh.has_text_frame and sh.top < I(1.2) and len(sh.text_frame.text) > 120:
            for para in sh.text_frame.paragraphs:
                for r in para.runs: r.font.size = Pt(20)
prs.save(sys.argv[2])
