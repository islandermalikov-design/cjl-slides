"""Увеличивает шрифт до 16 pt (где помещается) прямо в готовой презентации пользователя.
Правит: текстовые блоки, таблицы, шрифты в диаграммах; пересчитывает поля диаграмм и овалы вокруг основной группы.
Использование: python bump_fonts.py in.pptx out.pptx
"""
import sys, copy, re
from pptx import Presentation
from pptx.util import Pt, Emu
from lxml import etree

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
EMU = 914400.0
MIN = 16.0


def q(tag):
    pre, name = tag.split(":")
    return "{%s}%s" % (NS[pre], name)


# ---------------------------------------------------------------- текст
def bump_text_frame(tf, minimum=MIN, cap=None):
    for para in tf.paragraphs:
        for r in para.runs:
            sz = r.font.size.pt if r.font.size else None
            if sz is None or sz < minimum:
                r.font.size = Pt(minimum)
        # размер «по умолчанию» для конца абзаца
        epr = para._p.find(q("a:endParaRPr"))
        if epr is not None and epr.get("sz") and int(epr.get("sz")) < minimum * 100:
            epr.set("sz", str(int(minimum * 100)))


def is_tag(shape):
    return shape.has_text_frame and shape.text_frame.text.strip().startswith("Положение ") and len(shape.text_frame.text) < 14


# ---------------------------------------------------------------- диаграммы
def chart_info(cs):
    plot = cs.find(".//c:plotArea", NS)
    kinds = [e.tag.split("}")[1] for e in plot if e.tag.endswith("Chart")]
    kind = kinds[0] if kinds else ""
    bar_dir = plot.find(".//c:barDir", NS)
    horizontal = bar_dir is not None and bar_dir.get("val") == "bar"
    grouping = plot.find(".//c:grouping", NS)
    stacked = grouping is not None and grouping.get("val") in ("stacked", "percentStacked")
    nser = len(plot.findall(".//c:ser", NS))
    ncat = len(plot.find(".//c:ser", NS).findall(".//c:cat//c:pt", NS)) if nser else 0
    has_title = cs.find("./c:chart/c:title", NS) is not None
    has_legend = cs.find("./c:chart/c:legend", NS) is not None
    cat_title = plot.find("./c:catAx/c:title", NS) is not None
    val_ax = plot.find("./c:valAx", NS)
    val_title = val_ax is not None and val_ax.find("./c:title", NS) is not None
    val_hidden = val_ax is not None and val_ax.find("./c:delete", NS) is not None and val_ax.find("./c:delete", NS).get("val") == "1"
    return dict(kind=kind, horizontal=horizontal, stacked=stacked, nser=nser, ncat=ncat, has_title=has_title, has_legend=has_legend,
                cat_title=cat_title, val_title=val_title, val_hidden=val_hidden)


def set_sizes(root, sz):
    for el in root.iter(q("a:defRPr"), q("a:rPr")):
        el.set("sz", str(int(sz * 100)))


def bump_chart(chart, fx, fy, fw, fh):
    cs = chart._chartSpace
    info = chart_info(cs)
    wide = fw >= 5.8
    small_cat = info["ncat"] <= 3
    # шрифты
    t_sz = 16 if wide else 14
    leg_sz = 16 if wide else 14
    cat_sz = 16 if (wide or info["kind"] == "radarChart" or (not info["horizontal"] and info["ncat"] >= 6)) else 14
    if info["horizontal"]:
        cat_sz = 16 if wide else 14
    val_sz = 14
    ax_sz = 14
    data_sz = 14 if (small_cat and wide) else (12 if wide else 11)
    if info["horizontal"]:
        data_sz = 14 if wide else 12
        if info["nser"] >= 3:
            data_sz = 12
            leg_sz = 14
    # заголовок диаграммы: длинный -> 14 pt
    ct = cs.find("./c:chart/c:title", NS)
    if ct is not None:
        for a_t in ct.iter(q("a:t")):
            rep = {"Недостаточное обеспечение (НедВО)": "Недостаточное (НедВО)", "Нормальное обеспечение (НВО)": "Нормальное (НВО)", "Избыточное обеспечение (ИВО)": "Избыточное (ИВО)"}
            if a_t.text in rep:
                a_t.text = rep[a_t.text]
        ttxt = "".join(x.text or "" for x in ct.iter(q("a:t")))
        if len(ttxt) > fw * 0.75 * 9.0 / 1.0 * (1 if t_sz == 16 else 1.15):
            t_sz = 14
        if fw < 5 and len(ttxt) > 30:
            t_sz = 13
        set_sizes(ct, t_sz)
    lg = cs.find("./c:chart/c:legend", NS)
    if lg is not None:
        set_sizes(lg, leg_sz)
    plot = cs.find(".//c:plotArea", NS)
    for ax in plot.findall("./c:catAx", NS):
        t = ax.find("./c:title", NS)
        if t is not None:
            set_sizes(t, ax_sz)
        tx = ax.find("./c:txPr", NS)
        if tx is not None:
            set_sizes(tx, cat_sz)
    for ax in plot.findall("./c:valAx", NS):
        t = ax.find("./c:title", NS)
        if t is not None:
            set_sizes(t, ax_sz)
            for a_t in t.iter(q("a:t")):
                if a_t.text == "Баллы (шкала Ликерта)":
                    a_t.text = "Баллы"
        tx = ax.find("./c:txPr", NS)
        if tx is not None:
            set_sizes(tx, val_sz)
    for dl in plot.iter(q("c:dLbls")):
        tx = dl.find("./c:txPr", NS)
        if tx is not None:
            set_sizes(tx, data_sz)
        for d in dl.findall("./c:dLbl", NS):
            runs = d.findall(".//a:r", NS)
            for r in runs:
                rpr = r.find("./a:rPr", NS)
                txt = r.find("./a:t", NS).text or ""
                is_mark = txt.strip() in ("*", "▲", "*▲")
                rpr.set("sz", str(int((16 if is_mark else data_sz) * 100)))
    # поля области построения (дюймы -> доли)
    if info["kind"] != "radarChart":
        lay = plot.find("./c:layout/c:manualLayout", NS)
        if lay is not None:
            def g(tag):
                return float(lay.find("./c:" + tag, NS).get("val"))
            old_l = g("x") * fw
            if info["val_hidden"]:
                L = 0.45
            elif info["horizontal"]:
                L = 0.35 + max(old_l - 0.35, 0) * 1.28
                L = min(L, fw * 0.62)
            else:
                L = 1.05
            R = 0.2
            T = (0.55 if info["has_title"] else 0.15)
            legend_rows = 1
            if info["has_legend"]:
                names = ["".join(x.text or "" for x in sr.find("./c:tx", NS).iter(q("a:t"), q("c:v"))) for sr in plot.findall(".//c:ser", NS)]
                total = max(len(n) for n in names) * len(names) * (0.1 if leg_sz == 16 else 0.088) + 0.4 * len(names)
                if info["nser"] >= 4:
                    legend_rows = 2
                elif info["horizontal"] and info["nser"] == 3:
                    legend_rows = 3  # длинные названия комплексов лечения
                elif info["nser"] == 3 and not wide:
                    legend_rows = 2
                elif total > fw - 0.2:
                    legend_rows = 2
            leg_h = (0.12 + 0.31 * legend_rows) if info["has_legend"] else 0
            B = leg_h + (0.34 if cat_sz >= 16 else 0.3) + (0.32 if info["cat_title"] else 0) + (0.38 if info["stacked"] and info["ncat"] == 2 else 0)
            pw, ph = fw - L - R, fh - T - B
            vals = {"x": L / fw, "y": T / fh, "w": pw / fw, "h": ph / fh}
            for k, v in vals.items():
                lay.find("./c:" + k, NS).set("val", "%.5f" % v)
            return dict(plot=(fx + L, fy + T, pw, ph), ncat=info["ncat"], horizontal=info["horizontal"], kind=info["kind"])
    return dict(plot=None, ncat=info["ncat"], horizontal=info["horizontal"], kind=info["kind"])



# ---------------------------------------------------------------- точечные правки
def I(v):
    return Emu(int(round(v * EMU)))


def place(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)


def by_name(slide, name, startswith=False):
    return [sh for sh in slide.shapes if (sh.name.startswith(name) if startswith else sh.name == name)]


def find_slide(prs, prefix):
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix) and sh.top < I(1.2):
                return sl
    return None


def set_font(tf, pt):
    for para in tf.paragraphs:
        for r in para.runs:
            r.font.size = Pt(pt)


def set_align_left(tf):
    from pptx.enum.text import PP_ALIGN
    for para in tf.paragraphs:
        para.alignment = PP_ALIGN.LEFT


NOTE_MAP = [
    ("1 — контрольная группа", "1 — контрольная, 2 — сравнения, 3 — основная группа; а — СРК-З, б — СРК-Д"),
    ("В ячейках:", "В ячейках: до лечения, после лечения, p; Me [LQ; UQ]. VLF, LF, HF — мощность спектра (мс²); LF/HF — вагосимпатический баланс."),
]


def shorten_note(text):
    for k, v in NOTE_MAP:
        if text.startswith(k):
            return v
    if text.startswith("Примечание:"):
        if "до лечения и через 6 месяцев" in text:
            return "Примечание: * – различия внутри групп до лечения и через 6 мес.; ▲ – различия между группами через 6 мес. (p<0,05)"
        if "основной группы с контрольной и сравнения после лечения" in text:
            return "Примечание: * – различия внутри групп до и после лечения; ▲ – различия основной группы с контрольной и сравнения (p<0,05)"
        if text.startswith("Примечание: ▲"):
            if "через 6 месяцев" in text:
                return "Примечание: ▲ – различия основной группы с контрольной и сравнения через 6 мес. (p<0,05)"
            return "Примечание: ▲ – различия основной группы с контрольной и сравнения (p<0,05)"
        if "внутри групп до и после лечения" in text:
            return "Примечание: * – различия внутри групп до и после лечения; ▲ – различия между группами после лечения (p<0,05)"
    return text


def fix_notes(slide):
    is_radar = any(sh.chart._chartSpace.find(".//c:radarChart", NS) is not None for sh in slide.shapes if sh.has_chart)
    for sh in by_name(slide, "Примечание (достоверность)"):
        n = 0
        for para in sh.text_frame.paragraphs:
            if para.runs:
                full = "".join(r.text for r in para.runs)
                new = shorten_note(full)
                if new != full:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""
            n += 1
        lines = n
        if not is_radar:
            place(sh, y=7.03 - 0.29 * lines - 0.02, h=0.29 * lines + 0.04)


def fix_titles(slide):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.has_text_frame and "TITLE" in str(sh.placeholder_format.type):
            tf = sh.text_frame
            n = len(tf.text)
            size = 26 if n <= 56 else (24 if n <= 130 else 22)
            for para in tf.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(size)


def custom_fixes(prs):
    # --- 2. Актуальность
    s = find_slide(prs, "Актуальность")
    if s:
        for nm in ("Факт 1", "Факт 2", "Факт 3"):
            for sh in by_name(s, nm): place(sh, y=2.9, h=1.05)
        for sh in by_name(s, "Вывод по актуальности"): place(sh, y=4.05, h=1.0)
        for nm in ("Пустой блок 1", "Пустой блок 2"):
            for sh in by_name(s, nm):
                place(sh, y=5.2, h=2.1)
                set_align_left(sh.text_frame)
    # --- 5. Критерии
    s = find_slide(prs, "Критерии")
    if s:
        for sh in s.shapes:
            if sh.has_text_frame and "перечень" in sh.name:
                set_font(sh.text_frame, 14)
                place(sh, y=1.85, h=5.1)
            elif sh.has_text_frame and sh.name.startswith("Критерии"):
                place(sh, y=1.25, h=0.55)
    # --- 6. Дизайн
    s = find_slide(prs, "Дизайн исследования")
    if s:
        for sh in by_name(s, "Рандомизация"): place(sh, y=2.27, h=0.45)
        for sh in s.shapes:
            if sh.name.startswith("Группа ") or sh.name.startswith("Основная группа"):
                if sh.width / EMU < 4.1 and abs(sh.top / EMU - 3.05) < 0.05: place(sh, y=2.87, h=0.38)
            if sh.name.startswith("Подгруппа"): place(sh, y=3.25, h=0.62)
            if sh.name.startswith("Метод "): place(sh, y=4.07, h=1.95)
            if sh.name.startswith("Этап "): place(sh, y=6.12, h=0.9)
            if sh.name == "Стрелка":
                y = sh.top / EMU
                if abs(y - 2.15) < 0.03: place(sh, y=2.15, h=0.12)
                elif abs(y - 2.85) < 0.03: place(sh, y=2.72, h=0.15)
                elif abs(y - 4.15) < 0.03: place(sh, y=3.87, h=0.2)
                elif abs(y - 6.53) < 0.03: place(sh, y=6.57)
        for sh in by_name(s, "Метод 1"): set_font(sh.text_frame, 16)
    # --- 7. Методы исследования
    s = find_slide(prs, "Методы исследования")
    if s:
        top = {"Клинические": (0.5, 3.7), "Лабораторные": (4.35, 5.6), "Инструментальные": (10.1, 2.73)}
        bottom = {"Исследование вегетативной регуляции": (0.5, 3.7), "Исследование психоэмоционального статуса": (4.35, 5.6), "Исследование качества жизни": (10.1, 2.73)}
        for sh in s.shapes:
            for k, (x, w) in top.items():
                if sh.name == k: place(sh, x, 1.25, w, 0.6)
                elif sh.name == k + " — перечень": place(sh, x, 1.9, w, 2.95)
            for k, (x, w) in bottom.items():
                if sh.name == k: place(sh, x, 4.95, w, 0.6)
                elif sh.name == k + " — перечень": place(sh, x, 5.6, w, 1.4)
            if sh.has_text_frame and "перечень" in sh.name:
                set_align_left(sh.text_frame)
            if sh.has_text_frame and sh.name.startswith("Исследование") and "перечень" not in sh.name:
                set_font(sh.text_frame, 16)
    # --- 8. Методы лечения: таблица 14 pt
    s = find_slide(prs, "Методы лечения")
    if s:
        for sh in s.shapes:
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells: set_font(cell.text_frame, 14)
    # --- 11. Стул: таблица и обводка
    s = find_slide(prs, "Нормализация стула")
    if s:
        for sh in s.shapes:
            if sh.has_table:
                t = sh.table
                for i, wv in enumerate((2.5, 2.0, 2.13)): t.columns[i].width = I(wv)
                hs = [0.4, 0.5] + [0.6] * 6
                for i, hv in enumerate(hs): t.rows[i].height = I(hv)
                for row in t.rows:
                    for cell in row.cells: set_font(cell.text_frame, 15)
                top = sh.top / EMU
                for ov in by_name(s, "Овал: основная группа"):
                    place(ov, x=sh.left / EMU - 0.08, y=top + 0.9 + 4 * 0.6 - 0.06, w=sh.width / EMU + 0.16, h=1.2 + 0.12)
    # --- 15. Микробиоценоз: таблица
    s = find_slide(prs, "Динамика микробиоценоза")
    if s:
        for sh in s.shapes:
            if sh.has_table:
                t = sh.table
                widths = [3.3, 0.9, 0.95, 0.9, 0.95, 0.9, 1.0, 1.15, 1.15, 1.13]
                for i, wv in enumerate(widths): t.columns[i].width = I(wv)
                for row in t.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for r in para.runs:
                                txt = r.text
                                r.font.size = Pt(12 if txt.startswith("(p=") else 14)
                for i, row in enumerate(t.rows): row.height = I(0.4 if i == 0 else (0.34 if i == 1 else 0.55))
                x0 = sh.left / EMU
                for ov in by_name(s, "Овал: основная группа"):
                    place(ov, x=x0 + sum(widths[:5]) - 0.04, y=sh.top / EMU - 0.06, w=widths[5] + widths[6] + 0.08, h=0.4 + 0.34 + 8 * 0.55 + 0.1)
                for nt in s.shapes:
                    if nt.has_text_frame and nt.text_frame.text.startswith("до — до лечения"):
                        place(nt, y=6.6, h=0.4)
                        set_font(nt.text_frame, 14)
    # --- радары SF-36: уменьшить графики, увеличить текст примечаний
    for key in ("Качество жизни (SF-36) у пациентов с СРК-З", "Качество жизни (SF-36) у пациентов с СРК-Д"):
        s = find_slide(prs, key)
        if not s: continue
        for sh in s.shapes:
            if sh.has_chart:
                place(sh, y=1.2, h=3.85)
            if sh.name == "Овал: основная группа":
                place(sh, y=1.1, h=4.05)
            if sh.name == "Примечание (достоверность)":
                n = len(sh.text_frame.paragraphs)
                place(sh, y=5.2, h=1.75)
    # --- схема патогенеза: пояснения
    s = find_slide(prs, "Влияние лечебных комплексов")
    if s:
        for sh in by_name(s, "Пояснение АМП"): place(sh, x=10.1, y=1.95, w=2.73, h=1.5)
        for sh in by_name(s, "Пояснение КВЧ"): place(sh, x=9.6, y=3.42, w=3.23, h=1.1)
        for sh in s.shapes:
            if sh.has_text_frame and sh.name.startswith("Пояснение"):
                set_font(sh.text_frame, 16)
    # --- выводы: высоты карточек под 16 pt
    for key, hs in (("Выводы (1–3)", (2.1, 1.65, 2.05)), ("Выводы (4–6)", (1.65, 1.6, 2.45))):
        s = find_slide(prs, key)
        if not s: continue
        y = 1.0
        cards = sorted([sh for sh in s.shapes if sh.name.startswith("Вывод ")], key=lambda z: z.top)
        nums = sorted([sh for sh in s.shapes if sh.name.startswith("Номер вывода")], key=lambda z: z.top)
        for c, n, h in zip(cards, nums, hs):
            place(c, y=y, h=h)
            place(n, y=y + 0.08)
            y += h + 0.08
    # --- заголовки и подписи на всех слайдах
    for sl in prs.slides:
        fix_titles(sl)


# ---------------------------------------------------------------- основной проход
def process(src, dst):
    prs = Presentation(src)
    for idx, slide in enumerate(prs.slides, 1):
        charts = []
        ovals = []
        fix_notes(slide)
        note_top = min([n.top / EMU for n in by_name(slide, "Примечание (достоверность)")] or [99])
        for sh in slide.shapes:
            if sh.has_chart:
                if sh.top / EMU + sh.height / EMU > note_top - 0.02 and sh.top / EMU < note_top and not sh.name.startswith(("Контрольная", "Группа", "Основная")):
                    sh.height = I(note_top - 0.03 - sh.top / EMU)
                fx, fy, fw, fh = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
                geo = bump_chart(sh.chart, fx, fy, fw, fh)
                geo["frame"] = (fx, fy, fw, fh)
                charts.append(geo)
            elif sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        bump_text_frame(cell.text_frame)
            elif sh.has_text_frame:
                if sh.name.startswith("Овал"):
                    ovals.append(sh)
                    continue
                if is_tag(sh):
                    bump_text_frame(sh.text_frame, 16)
                    sh.top, sh.height = I(0.04), I(0.32)
                    continue
                if sh.is_placeholder and sh.placeholder_format.type is not None and "SLIDE_NUMBER" in str(sh.placeholder_format.type):
                    continue
                bump_text_frame(sh.text_frame)
        # овалы вокруг основной группы — пересчёт под новые поля диаграммы
        for ov in ovals:
            cx, cy = (ov.left + ov.width / 2) / EMU, (ov.top + ov.height / 2) / EMU
            for geo in charts:
                fx, fy, fw, fh = geo["frame"]
                if fx <= cx <= fx + fw and fy <= cy <= fy + fh and geo["plot"] and not geo["horizontal"]:
                    n = {6: 2, 3: 1}.get(geo["ncat"])
                    if not n:
                        break
                    px, py, pw, ph = geo["plot"]
                    tot = geo["ncat"]
                    pad = 0.06
                    x = px + pw * (tot - n) / tot - pad
                    w = min(pw * n / tot + 2 * pad, fx + fw - x - 0.02)
                    ov.left, ov.top, ov.width, ov.height = Emu(int(x * EMU)), Emu(int((py - 0.02) * EMU)), Emu(int(w * EMU)), Emu(int((ph + 0.14) * EMU))
                    break
    custom_fixes(prs)
    prs.save(dst)


if __name__ == "__main__":
    process(sys.argv[1], sys.argv[2])
