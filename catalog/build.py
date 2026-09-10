# -*- coding: utf-8 -*-
"""Сборка каталога Irbis: python3 build.py  →  index.html + Каталог_Irbis.xlsx"""
import html
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from data import CATEGORIES, PRODUCTS          # noqa: E402
from svg_art import render as art              # noqa: E402

e = html.escape


def is_tbd(row):
    return len(row) > 2 and row[2] == "tbd"


def cell(row, cls):
    val = row[1]
    if is_tbd(row):
        if val in ("—", "", None):
            return '<dd class="%s tbd">уточняется</dd>' % cls
        return '<dd class="%s tbd">%s<span class="q" title="Требует подтверждения">?</span></dd>' % (cls, e(val))
    return '<dd class="%s">%s</dd>' % (cls, e(val))


def card(p):
    cat = p["cat"]
    hay = " ".join([p["name"], p.get("analog", ""), p.get("sub", "")] +
                   ["%s %s" % (r[0], r[1]) for r in p["key"]] +
                   ["%s %s" % (r[0], r[1]) for r in p.get("specs", [])]).lower()

    keys = "".join('<div class="kv"><dt>%s</dt>%s</div>' % (e(r[0]), cell(r, "v"))
                   for r in p["key"])

    more = ""
    if p.get("specs"):
        rows = "".join('<div class="kv2"><dt>%s</dt>%s</div>' % (e(r[0]), cell(r, "v2"))
                       for r in p["specs"])
        more = ('<details class="more"><summary>Полные характеристики'
                '<span class="cnt">%d</span></summary><dl>%s</dl></details>' % (len(p["specs"]), rows))

    note = '<p class="note">%s</p>' % e(p["note"]) if p.get("note") else ""
    src = ('<a class="src" href="%s" target="_blank" rel="noopener">Карточка поставщика</a>' % e(p["src"])
           if p.get("src") else '<span class="src--none">Ссылка в исходном файле не указана</span>')
    sub = '<p class="sub">%s</p>' % e(p["sub"]) if p.get("sub") else ""
    analog = ('<p class="analog"><span>Платформа</span>%s</p>' % e(p["analog"])
              if p.get("analog") else "")

    return ('<article class="card" data-cat="%s" data-q="%s">\n'
            '<figure class="shot">%s<img src="photos/%s.jpg" alt="%s" loading="lazy" onerror="this.remove()"></figure>\n'
            '<div class="body"><div class="titles"><h3>%s</h3>%s</div>%s'
            '<dl class="keys">%s</dl>%s%s<div class="foot">%s</div></div></article>'
            % (cat, e(hay), art(cat), p["id"], e(p["name"]),
               e(p["name"]), sub, analog, keys, more, note, src))


SEARCH_ICON = ('<svg width="15" height="15" viewBox="0 0 16 16" fill="none" stroke="currentColor" '
               'stroke-width="1.7" aria-hidden="true"><circle cx="7" cy="7" r="4.6"/>'
               '<path d="M10.4 10.4 14 14"/></svg>')


def build():
    total = sum(len(p["key"]) for p in PRODUCTS)
    tbd = sum(1 for p in PRODUCTS for r in p["key"] if is_tbd(r))
    filled = round(100 * (total - tbd) / total)
    flagged = [p for p in PRODUCTS if p.get("note")]

    chips = "".join(
        '<button class="chip" type="button" aria-pressed="false" data-filter="%s">%s<span>%d</span></button>'
        % (c["id"], e(c["title"]), sum(1 for p in PRODUCTS if p["cat"] == c["id"]))
        for c in CATEGORIES)

    sections = []
    for c in CATEGORIES:
        items = [p for p in PRODUCTS if p["cat"] == c["id"]]
        sections.append(
            '<section class="cat" id="%s"><div class="cat-head"><h2>%s</h2>'
            '<p class="cat-lead">%s</p>'
            '<p class="cat-keys"><span>Ключевые параметры</span>%s</p>'
            '<p class="cat-count">%d</p></div><div class="grid">%s</div></section>'
            % (c["id"], e(c["title"]), e(c["lead"]), e(" · ".join(c["keys"])),
               len(items), "".join(card(p) for p in items)))

    ledger = "".join('<li><b>%s</b><span>%s</span></li>' % (e(p["name"]), e(p["note"]))
                     for p in flagged)

    css = open(os.path.join(HERE, "style.css"), encoding="utf-8").read()
    js = open(os.path.join(HERE, "app.js"), encoding="utf-8").read()

    head = (
        '<meta charset="utf-8">\n<title>Каталог Irbis</title>\n'
        '<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>\n'
        '<link rel="stylesheet" href="https://fonts.googleapis.com/css2?'
        'family=Oswald:wght@400;500;600&family=PT+Sans:wght@400;700&'
        'family=IBM+Plex+Mono:wght@400;500&display=swap">\n'
        '<style>\n%s\n</style>\n' % css)

    body = (
        '<div class="page">\n'
        '<header class="masthead"><div><p class="brand">Irbis · спецтехника</p>'
        '<h1>Каталог<br>техники</h1>'
        '<p class="lede">Семь товарных групп, %d моделей. У каждой позиции — ключевые параметры '
        'группы, полная спецификация и ссылка на карточку поставщика, по которой велась сверка.</p></div>'
        '<dl class="tally">'
        '<div><dt>Группы</dt><dd>%d</dd></div>'
        '<div><dt>Модели</dt><dd>%d</dd></div>'
        '<div><dt>Готовность данных</dt><dd>%d%%</dd></div>'
        '<div class="warn"><dt>К уточнению</dt><dd>%d</dd></div>'
        '</dl></header>\n'
        '<nav class="toolbar" aria-label="Фильтры каталога">'
        '<label class="search">%s<input id="q" type="search" placeholder="Поиск по моделям, платформам и параметрам" '
        'autocomplete="off" aria-label="Поиск по каталогу"></label>'
        '<div class="chips">%s</div></nav>\n'
        '<main>%s<p class="empty" id="empty" hidden>Ничего не найдено. Измените запрос или снимите фильтр.</p></main>\n'
        '<section class="ledger"><h2>Расхождения и открытые вопросы</h2>'
        '<p>Позиции, где данные исходной таблицы расходятся с каталогами поставщиков либо '
        'отсутствуют в открытых источниках. До подтверждения не выносить в коммерческое предложение.</p>'
        '<ol>%s</ol></section>\n'
        '<footer class="colophon">'
        '<p><b>Фото</b><br>Положите снимок в <code>photos/&lt;код&gt;.jpg</code> — карточка подхватит его '
        'автоматически, вместо контурной схемы. Коды перечислены в <code>photos/README.md</code>.</p>'
        '<p><b>Данные</b><br>Правятся в <code>data.py</code>, страница пересобирается '
        'командой <code>python3 build.py</code>.</p>'
        '<p><b>Курсивом и знаком «?»</b><br>отмечены значения, требующие подтверждения '
        'у поставщика: %d из %d ключевых параметров.</p>'
        '</footer>\n</div>\n<script>\n%s\n</script>\n'
        % (len(PRODUCTS), len(CATEGORIES), len(PRODUCTS), filled, len(flagged),
           SEARCH_ICON, chips, "\n".join(sections), ledger, tbd, total, js))

    return head + body


def xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        print("  openpyxl не установлен — xlsx пропущен")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Каталог"
    head = ["№", "Группа", "Модель Irbis", "Базовая платформа", "Исполнение",
            "Ключевой параметр 1", "Значение 1", "Ключевой параметр 2", "Значение 2",
            "Ключевой параметр 3", "Значение 3", "Прочие характеристики",
            "Источник", "Код фото", "Примечание"]
    ws.append(head)
    thin = Side(style="thin", color="D0D0D0")
    for i, c in enumerate(ws[1], 1):
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="2B4757")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    cat_name = {c["id"]: c["title"] for c in CATEGORIES}
    n = 0
    for c in CATEGORIES:
        for p in [x for x in PRODUCTS if x["cat"] == c["id"]]:
            n += 1
            k = p["key"]
            def v(r):
                return ("уточняется" if r[1] in ("—", "", None) else r[1] + " (уточнить)") if is_tbd(r) else r[1]
            other = "; ".join("%s: %s" % (r[0], v(r)) for r in p.get("specs", []))
            ws.append([n, cat_name[p["cat"]], p["name"], p.get("analog", ""), p.get("sub", ""),
                       k[0][0], v(k[0]), k[1][0], v(k[1]), k[2][0], v(k[2]),
                       other, p.get("src", ""), "photos/%s.jpg" % p["id"], p.get("note", "")])
    widths = [4, 26, 20, 22, 26, 24, 18, 24, 18, 24, 18, 60, 44, 22, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.border = Border(bottom=thin)
            if c.row > 1:
                c.alignment = Alignment(vertical="top", wrap_text=c.column in (12, 15))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:O%d" % (n + 1)
    out = os.path.join(HERE, "Каталог_Irbis.xlsx")
    wb.save(out)
    print("  →", out)


if __name__ == "__main__":
    page = build()
    out = os.path.join(HERE, "index.html")
    with open(out, "w", encoding="utf-8") as f:
        f.write(page)
    print("  →", out, "(%.1f КБ)" % (len(page.encode()) / 1024))
    xlsx()
