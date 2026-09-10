# -*- coding: utf-8 -*-
"""
Design system for the MSEO round-table deck.

Palette harmonises two brands:
  * Форум риск-менеджеров 2026 (рисковики.рф) — deep violet hero gradient, fuchsia accent
  * МСЭО — navy wordmark, teal check-mark
"""

import os
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- palette ---
VIOLET_DEEP = RGBColor(0x24, 0x0D, 0x4E)   # hero gradient start
VIOLET      = RGBColor(0x3A, 0x14, 0x74)
VIOLET_MID  = RGBColor(0x5B, 0x21, 0xA8)   # hero gradient end
VIOLET_SOFT = RGBColor(0x4A, 0x22, 0x8C)
MAGENTA     = RGBColor(0xE2, 0x1C, 0xA8)   # site CTA / accent
MAGENTA_HI  = RGBColor(0xFF, 0x5C, 0xD2)
TEAL        = RGBColor(0x2A, 0x8D, 0xB6)   # МСЭО mark
TEAL_DEEP   = RGBColor(0x1B, 0x6A, 0x8C)
NAVY        = RGBColor(0x26, 0x2E, 0x41)   # МСЭО wordmark / body text
SLATE       = RGBColor(0x5E, 0x67, 0x7D)
SLATE_LIGHT = RGBColor(0x8A, 0x92, 0xA6)
LILAC       = RGBColor(0xF6, 0xF2, 0xFC)   # light page ground
LILAC_2     = RGBColor(0xEA, 0xE1, 0xF7)
LILAC_3     = RGBColor(0xDA, 0xCD, 0xEE)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
INK         = RGBColor(0x1A, 0x1A, 0x1E)

# ------------------------------------------------------------------ fonts ---
F_HEAD = "Segoe UI Semibold"   # slide titles, metric numbers
F_BOLD = "Segoe UI"            # body, labels (bold applied per-run)
F_SER  = "Georgia"             # moderator questions / pull quotes

# ---------------------------------------------------------------- geometry --
SW, SH = Inches(13.3333), Inches(7.5)
M      = Inches(0.68)          # side margin
CW     = SW - 2 * M            # content width
TITLE_T = Inches(0.62)
BODY_T  = Inches(1.72)
FOOT_Y  = Inches(6.86)

ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
LOGO_COLOR = os.path.join(ASSETS, "logo_color.png")
LOGO_WHITE = os.path.join(ASSETS, "logo_white.png")
MARK_COLOR = os.path.join(ASSETS, "mark_color.png")
MARK_WHITE = os.path.join(ASSETS, "mark_white.png")

A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


# ------------------------------------------------------------- xml helpers --
def _sub(parent, tag, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def set_alpha(shape, alpha, line=False):
    """Apply transparency (0..1 = opacity) to a shape's solid fill or line."""
    spPr = shape._element.spPr
    holder = spPr.find(qn("a:ln")) if line else spPr
    if holder is None:
        return
    solid = holder.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    _sub(srgb, "a:alpha", val=int(alpha * 100000))


def set_spacing(run, pts):
    """Letter-spacing in points (can be negative)."""
    run.font._rPr.set("spc", str(int(pts * 100)))


def no_line(shape):
    shape.line.fill.background()


def soft_shadow(shape, blur=18, dist=6, alpha=0.10, color="1B1040"):
    """Outer shadow via raw DrawingML (python-pptx has no high-level API)."""
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    xml = (
        '<a:effectLst %s>'
        '<a:outerShdw blurRad="%d" dist="%d" dir="5400000" rotWithShape="0">'
        '<a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr>'
        '</a:outerShdw></a:effectLst>'
        % (A_NS, Pt(blur), Pt(dist), color, int(alpha * 100000))
    )
    spPr.append(etree.fromstring(xml))


# ------------------------------------------------------------ shape helpers --
def rect(slide, l, t, w, h, color, shape=MSO_SHAPE.RECTANGLE, radius=None,
         alpha=None, rot=None):
    s = slide.shapes.add_shape(shape, int(l), int(t), int(w), int(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    no_line(s)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if alpha is not None:
        set_alpha(s, alpha)
    if rot is not None:
        s.rotation = rot
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def ring(slide, cx, cy, d, color, width_pt=1.4, alpha=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(cx - d / 2), int(cy - d / 2),
                               int(d), int(d))
    s.fill.background()
    s.shadow.inherit = False
    s.line.color.rgb = color
    s.line.width = Pt(width_pt)
    if alpha is not None:
        set_alpha(s, alpha, line=True)
    return s


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(int(l), int(t), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tb, tf


def para(tf, first, text, size, color, font=F_BOLD, bold=False, italic=False,
         align=PP_ALIGN.LEFT, spacing=None, line=1.15, before=0, after=0,
         caps=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if caps:
        text = text.upper()
    chunks = text.split("\n")
    for i, chunk in enumerate(chunks):
        if i:
            p._p.append(etree.Element(qn("a:br")))
        r = p.add_run()
        r.text = chunk
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color
        if spacing:
            set_spacing(r, spacing)
        # make the East-Asian / complex-script slots follow the latin typeface
        # so Cyrillic never silently falls back to a different family
        rPr = r.font._rPr
        for tag in ("a:ea", "a:cs"):
            for old in rPr.findall(qn(tag)):
                rPr.remove(old)
            _sub(rPr, tag, typeface=font)
    return p


def block(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of dicts consumed by para()."""
    _, tf = textbox(slide, l, t, w, h, anchor)
    for i, kw in enumerate(lines):
        para(tf, i == 0, **kw)
    return tf


def picture(slide, path, l, t, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = int(w)
    if h:
        kw["height"] = int(h)
    return slide.shapes.add_picture(path, int(l), int(t), **kw)


# ------------------------------------------------------------- backgrounds --
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg_light(slide, color=LILAC):
    rect(slide, 0, 0, SW, SH, color)


def bg_gradient(slide, c1=VIOLET_DEEP, c2=VIOLET_MID, angle=35):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, int(SW), int(SH))
    no_line(s)
    s.shadow.inherit = False
    f = s.fill
    f.gradient()
    stops = f.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    while len(stops._gsLst) > 2:                      # drop preset extra stops
        stops._gsLst.remove(stops._gsLst[-1])
    f.gradient_angle = angle
    return s


def hero_rings(slide, cx, cy, sizes=(7.4, 5.9, 4.4, 3.0), color=None,
               alpha=0.16, width=1.6):
    """Concentric outlines echoing the МСЭО check-mark circle + site hero."""
    color = color or MAGENTA_HI
    for i, d in enumerate(sizes):
        ring(slide, cx, cy, Inches(d), color,
             width_pt=width, alpha=alpha * (1 - i * 0.12))


# ----------------------------------------------------------- page furniture --
def eyebrow(slide, text, color=MAGENTA, t=None, l=None, size=10.5):
    t = TITLE_T if t is None else t
    _, tf = textbox(slide, l if l is not None else M, t, CW, Inches(0.26))
    para(tf, True, text=text, size=size, color=color, font=F_BOLD, bold=True,
         caps=True, spacing=1.7)


def title(slide, text, color=NAVY, t=None, size=27, w=None, l=None,
          font=F_HEAD, line=1.06):
    t = Inches(0.95) if t is None else t
    _, tf = textbox(slide, l if l is not None else M, t, w or CW, Inches(1.15))
    para(tf, True, text=text, size=size, color=color, font=font, bold=True,
         line=line)
    return tf


def footer(slide, source=None, page=None, dark=False, partner=True, fw=None):
    line_c = LILAC_3 if not dark else VIOLET_SOFT
    fw = CW if fw is None else fw
    rect(slide, M, FOOT_Y, fw, Emu(9525), line_c)
    if source:
        _, tf = textbox(slide, M, FOOT_Y + Inches(0.13), fw - Inches(1.6),
                        Inches(0.4))
        para(tf, True, text=source, size=8.5,
             color=SLATE_LIGHT if not dark else RGBColor(0x9C, 0x8C, 0xC4),
             font=F_BOLD, line=1.2)
    if partner:
        picture(slide, MARK_WHITE if dark else MARK_COLOR,
                M + fw - Inches(1.02), FOOT_Y + Inches(0.14), h=Inches(0.28))
    if page is not None:
        _, tf = textbox(slide, M + fw - Inches(0.60), FOOT_Y + Inches(0.16),
                        Inches(0.60), Inches(0.3))
        para(tf, True, text=str(page), size=10,
             color=SLATE_LIGHT if not dark else RGBColor(0xB0, 0x9E, 0xD8),
             font=F_BOLD, bold=True, align=PP_ALIGN.RIGHT)


# ============================================================================
#  Приёмы из референс-листа: пончики, пилюли-теги, изображения, затемнения
# ============================================================================

IMG_CHAIN = os.path.join(ASSETS, "img_chain.jpg")
IMG_LAYERS = os.path.join(ASSETS, "img_layers.jpg")
IMG_INDUSTRIES = os.path.join(ASSETS, "img_industries.jpg")


def donut(slide, cx, cy, d, segments, thickness=0.30, start=270.0, gap=1.2):
    """Кольцевая диаграмма из дуг (blockArc).

    segments: [(value, color), ...]; start — угол начала (270° = 12 часов).
    """
    total = float(sum(v for v, _ in segments)) or 1.0
    a = start
    out = []
    for value, color in segments:
        sweep = 360.0 * value / total
        s = slide.shapes.add_shape(MSO_SHAPE.BLOCK_ARC,
                                   int(cx - d / 2), int(cy - d / 2),
                                   int(d), int(d))
        s.adjustments[0] = ((a + gap / 2) % 360.0) * 0.6
        s.adjustments[1] = ((a + sweep - gap / 2) % 360.0) * 0.6
        s.adjustments[2] = thickness
        s.fill.solid()
        s.fill.fore_color.rgb = color
        no_line(s)
        s.shadow.inherit = False
        out.append(s)
        a += sweep
    return out


def pill(slide, l, t, text, fill=None, text_color=None, size=8.5,
         border=None, h=None):
    """Тег-пилюля из паттерна «распределение рисков»; возвращает ширину."""
    h = h or Inches(0.235)
    w = Inches(0.28) + Inches(0.078) * len(text)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(l), int(t),
                               int(w), int(h))
    s.adjustments[0] = 0.5
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if border is None:
        no_line(s)
    else:
        s.line.color.rgb = border
        s.line.width = Pt(0.75)
    _, tf = textbox(slide, l, t + Inches(0.028), w, h)
    para(tf, True, text=text, size=size, color=text_color or WHITE,
         font=F_BOLD, bold=True, align=PP_ALIGN.CENTER, caps=True,
         spacing=0.4)
    return w


def pill_row(slide, l, t, texts, max_w, fill=None, text_color=None,
             gap=None, line_h=None, **kw):
    """Раскладывает пилюли в строки, перенося по ширине max_w."""
    gap = gap or Inches(0.09)
    line_h = line_h or Inches(0.32)
    x, y = l, t
    for txt in texts:
        w = Inches(0.28) + Inches(0.078) * len(txt)
        if x > l and x + w > l + max_w:
            x, y = l, y + line_h
        pill(slide, x, y, txt, fill=fill, text_color=text_color, **kw)
        x += w + gap
    return y + line_h


def image_fill(slide, path, l, t, w, h):
    """Вставляет картинку «под обрез»: заполняет бокс, лишнее обрезается."""
    from PIL import Image
    iw, ih = Image.open(path).size
    box_ar, img_ar = float(w) / float(h), float(iw) / float(ih)
    pic = slide.shapes.add_picture(path, int(l), int(t), width=int(w),
                                   height=int(h))
    if img_ar > box_ar:
        c = (1.0 - box_ar / img_ar) / 2.0
        pic.crop_left = pic.crop_right = c
    elif img_ar < box_ar:
        c = (1.0 - img_ar / box_ar) / 2.0
        pic.crop_top = pic.crop_bottom = c
    return pic


def scrim(slide, l, t, w, h, color=None, a_from=0.92, a_to=0.0, angle=0):
    """Градиентная «вуаль» поверх картинки — чтобы текст читался."""
    color = color or VIOLET_DEEP
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(l), int(t), int(w),
                               int(h))
    no_line(s)
    s.shadow.inherit = False
    f = s.fill
    f.gradient()
    stops = f.gradient_stops
    while len(stops._gsLst) > 2:
        stops._gsLst.remove(stops._gsLst[-1])
    stops[0].color.rgb = color
    stops[0].position = 0.0
    stops[1].color.rgb = color
    stops[1].position = 1.0
    f.gradient_angle = angle
    for i, alpha in ((0, a_from), (1, a_to)):
        srgb = stops._gsLst[i].find(qn("a:srgbClr"))
        for old in srgb.findall(qn("a:alpha")):
            srgb.remove(old)
        _sub(srgb, "a:alpha", val=int(alpha * 100000))
    return s


def tint(slide, l, t, w, h, color=None, alpha=0.42):
    """Ровная подложка поверх картинки."""
    return rect(slide, l, t, w, h, color or VIOLET_DEEP, alpha=alpha)
