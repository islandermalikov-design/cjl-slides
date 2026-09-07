#!/usr/bin/env python3
"""Собирает PPTX с редактируемым текстом из deck.json (геометрия снята с HTML)."""
import json, re, base64, os, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

PX = 7620                      # EMU на 1 px при слайде 1600x900 -> 13.333x7.5"
PT = PX / 12700.0              # 1 px в пунктах при этом масштабе (= 0.6)
Y_SHIFT = 0.0                  # общая правка вертикали (px), подбирается по рендеру
data = json.load(open('deck.json'))

def col(c):
    if not c: return None
    m = re.match(r'#([0-9a-fA-F]{6})', c)
    if m: return RGBColor.from_string(m.group(1).upper())
    m = re.findall(r'[\d.]+', c or '')
    if len(m) >= 3:
        if len(m) > 3 and float(m[3]) == 0: return None
        return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))
    return None

def font_for(family, weight):
    fam = (family or '').split(',')[0].strip().strip("'\"")
    if 'Playfair' in fam:
        return ('Playfair Display Black', False) if weight >= 800 else ('Playfair Display', True)
    if weight >= 700: return ('Noto Sans', True)
    if weight >= 600: return ('Noto Sans SemiBold', False)
    return ('Noto Sans', False)

ALIGN = {'left': PP_ALIGN.LEFT, 'start': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT,
         'end': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER, 'justify': PP_ALIGN.JUSTIFY}

os.makedirs('imgout', exist_ok=True)
img_files = {}
def img_path(src):
    if src in img_files: return img_files[src]
    head, b64 = src.split(',', 1)
    ext = 'png' if 'png' in head else 'jpg'
    p = f'imgout/img{len(img_files)}.{ext}'
    open(p, 'wb').write(base64.b64decode(b64))
    img_files[src] = p
    return p


def flatten(shp):
    """Убрать ссылку на стиль темы и любые эффекты (тень)."""
    sp = shp._element
    st = sp.find(qn('p:style'))
    if st is not None: sp.remove(st)
    shp.shadow.inherit = False

prs = Presentation()
prs.slide_width = Emu(1600 * PX)
prs.slide_height = Emu(900 * PX)
blank = prs.slide_layouts[6]

for sl in data:
    s = prs.slides.add_slide(blank)
    # белый фон
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF); bg.line.fill.background()
    flatten(bg)

    for sh in sl['shapes']:
        w, h = max(sh['w'], 1), max(sh['h'], 1)
        radius = sh.get('radius') or 0
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 1.5 else MSO_SHAPE.RECTANGLE
        shp = s.shapes.add_shape(kind, Emu(int(sh['x']*PX)), Emu(int(sh['y']*PX)),
                                 Emu(int(w*PX)), Emu(int(h*PX)))
        if radius > 1.5:
            try: shp.adjustments[0] = min(0.5, radius / min(w, h))
            except Exception: pass
        f = col(sh.get('fill'))
        if f is not None: shp.fill.solid(); shp.fill.fore_color.rgb = f
        else: shp.fill.background()
        ln = sh.get('line')
        if ln and col(ln['color']) is not None:
            shp.line.color.rgb = col(ln['color']); shp.line.width = Pt(max(ln['w'], 1) * PT)
            if ln.get('style') == 'dashed':
                shp.line._get_or_add_ln().append(shp.line._get_or_add_ln().makeelement(qn('a:prstDash'), {'val': 'dash'}))
        else:
            shp.line.fill.background()
        flatten(shp)

    for im in sl['images']:
        s.shapes.add_picture(img_path(im['full']), Emu(int(im['x']*PX)), Emu(int(im['y']*PX)),
                             Emu(int(im['w']*PX)), Emu(int(im['h']*PX)))

    for tx in sl['texts']:
        paras = [[]]
        for rn in tx['runs']:
            if rn.get('br'): paras.append([])
            else: paras[-1].append(rn)
        # если в браузере каждая строка = отдельный абзац (перенос по <br>),
        # то автоперенос не нужен — иначе PowerPoint ломает строки по-своему
        nowrap = tx.get('lines', 1) <= len(paras)
        w = tx['w'] + 4 if nowrap else max(tx.get('avail', tx['w']), tx['w'] * 1.06) + 4
        x = tx['x']
        align = tx.get('align')
        if align in ('right', 'end'):
            x = tx['x'] + tx['w'] - w            # правый край держим на месте
        elif align == 'center':
            x = tx['x'] - (w - tx['w']) / 2
        box = s.shapes.add_textbox(Emu(int(x*PX)), Emu(int((tx['y'] + Y_SHIFT)*PX)),
                                   Emu(int(w*PX)), Emu(int(max(tx['h'], 10)*PX)))
        tf = box.text_frame
        tf.word_wrap = not nowrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        # блок центрируем по вертикали: так первая строка не зависит от того,
        # как конкретный редактор трактует межстрочный интервал
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = 0; tf.margin_right = 0
        tf.margin_top = 0; tf.margin_bottom = 0
        first = True
        for pr in paras:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = ALIGN.get(tx.get('align'), PP_ALIGN.LEFT)
            p.line_spacing = Pt(tx['lh'] * PT)
            p.space_before = Pt(0); p.space_after = Pt(0)
            for rn in pr:
                t = rn['t']
                if rn.get('tt') == 'uppercase': t = t.upper()
                run = p.add_run(); run.text = t
                fam, bold = font_for(rn.get('family'), rn.get('weight') or 400)
                run.font.name = fam
                run.font.bold = bold
                run.font.size = Pt(rn['size'] * PT)
                c = col(rn.get('color'))
                if c is not None: run.font.color.rgb = c
                rPr = run._r.get_or_add_rPr()
                rPr.set('lang', 'ru-RU')
                if rn.get('ls'):
                    rPr.set('spc', str(int(rn['ls'] * PT * 100)))
                # кириллица должна брать тот же шрифт, что и латиница
                for tag in ('a:cs', 'a:ea'):
                    e = rPr.makeelement(qn(tag), {'typeface': fam}); rPr.append(e)

out = sys.argv[1] if len(sys.argv) > 1 else 'text.pptx'
prs.save(out)
print('saved', out, os.path.getsize(out)//1024, 'KB, slides:', len(prs.slides._sldIdLst))
