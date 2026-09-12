"""Spec -> native PowerPoint.

Every card is a real rounded rectangle, every rule a real line and every string
a real text frame, so the deck can be edited in PowerPoint (or imported into
Keynote / Google Slides) without touching an image editor.

Type is set in Arial: it is present on every machine and carries Cyrillic. To
switch to the corporate face, change PPTX_FONT below and rebuild.
"""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

import metrics
import spec
from spec import Img, Line, Rect, Text

HERE = os.path.dirname(os.path.abspath(__file__))

PPTX_FONT = "Arial"          # swap for "Alfa Sans" / "Styrene A" if installed
PX = 914400 / 144            # 1920 px across a 13.333 in slide -> EMU per px
BOLD_AT = 600                # weights at or above this become Arial Bold


def emu(px: float) -> int:
    return int(round(px * PX))


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip("#").upper())


def _add_rect(slide, el: Rect):
    if el.r:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, emu(el.x), emu(el.y), emu(el.w), emu(el.h))
        # roundRect adjustment is the corner radius as a share of the short side
        shape.adjustments[0] = el.r / min(el.w, el.h)
    else:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, emu(el.x), emu(el.y), emu(el.w), emu(el.h))
    shape.shadow.inherit = False
    if el.fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(el.fill)
    else:
        shape.fill.background()
    if el.stroke:
        shape.line.color.rgb = rgb(el.stroke)
        shape.line.width = Pt(el.sw * 0.5)
    else:
        shape.line.fill.background()
    shape.text_frame.word_wrap = False
    return shape


def _add_line(slide, el: Line):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, emu(el.x1), emu(el.y1), emu(el.x2), emu(el.y2))
    conn.line.color.rgb = rgb(el.color)
    conn.line.width = Pt(el.sw * 0.5)
    return conn


def _add_text(slide, el: Text):
    """One text frame per line: the baseline then lands exactly where the SVG
    puts it, whatever line-spacing model the reader applies."""
    asc = metrics.ascent("arial")
    for i, line in enumerate(el.lines):
        base = el.y + el.step * i
        size = max(r.size for r in line)
        width = sum(metrics.text_width(r.text, r.size, r.weight, r.tracking, "arial")
                    for r in line)
        top = base - size * asc
        box = slide.shapes.add_textbox(
            emu(el.x), emu(top), emu(width + size * 0.6), emu(size * 1.4))
        tf = box.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        for run_spec in line:
            run = para.add_run()
            run.text = run_spec.text
            font = run.font
            font.name = PPTX_FONT
            font.size = Pt(run_spec.size * 0.5)
            font.bold = run_spec.weight >= BOLD_AT
            font.color.rgb = rgb(run_spec.color)
            if run_spec.tracking:
                # spc is letter-spacing in 1/100 pt
                run.font._rPr.set("spc", str(int(run_spec.tracking * run_spec.size * 50)))


def _add_image(slide, el: Img):
    path = os.path.join(HERE, "assets", f"{el.name}.png")
    return slide.shapes.add_picture(path, emu(el.x), emu(el.y), emu(el.w), emu(el.h))


def build(slides, path: str) -> str:
    prs = Presentation()
    prs.slide_width = emu(spec.W)
    prs.slide_height = emu(spec.H)
    blank = prs.slide_layouts[6]

    for s in slides:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb("#FFFFFF")
        for el in s.elements:
            if isinstance(el, Rect):
                _add_rect(slide, el)
            elif isinstance(el, Line):
                _add_line(slide, el)
            elif isinstance(el, Text):
                _add_text(slide, el)
            elif isinstance(el, Img):
                _add_image(slide, el)
            else:
                raise TypeError(f"unsupported primitive: {el!r}")

    prs.save(path)
    return path
