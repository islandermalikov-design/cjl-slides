# -*- coding: utf-8 -*-
"""PPTX для направления E — Full-bleed Editorial. Геометрия зеркалит build_e.py 1:1."""
import pathlib, re, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

sys.path.insert(0, "build")
import deck_spec as S

ROOT = pathlib.Path(".").resolve()
OUT = ROOT / "restavraciya-obektov-E.pptx"

U_PX = 19.2                       # 1 "u" = 1% ширины эталонного слайда 1920px
PX_EMU = 6350
U_EMU = U_PX * PX_EMU             # EMU на 1 u
U_PT = U_PX * 0.5                 # pt на 1 u (px->pt при 1920px = 960pt)
SLIDE_W, SLIDE_H = Emu(int(100 * U_EMU)), Emu(int(56.25 * U_EMU))

FONT_SERIF = "Cormorant Garamond Light"
FONT_SANS = "Noto Sans"
FONT_SANS_SB = "Noto Sans SemiBold"

C_PRIMARY = RGBColor(0x1A, 0x18, 0x15)
C_SECONDARY = RGBColor(0xF2, 0xEC, 0xE3)
C_ACCENT = RGBColor(0x9A, 0x7F, 0x5E)
C_STONE = RGBColor(0xA9, 0xA0, 0x92)
C_TAUPE = RGBColor(0x6E, 0x64, 0x59)
C_LINE_L = ("26241D", 22)         # (hex, alpha%) — линия на светлом
C_LINE_D = ("F2ECE3", 18)         # линия на тёмном

PLACEHOLDER_SLOTS = {"borby-after", "nikoloyamskaya-before", "pochtovaya-after"}


def U(v):
    return Emu(int(round(v * U_EMU)))


def PT(v):
    return Pt(round(v * U_PT, 1))


def name_shape(shape, name, descr=None):
    cNvPr = shape._element.find(".//" + qn("p:cNvPr"))
    cNvPr.set("name", name)
    if descr:
        cNvPr.set("descr", descr)


def no_shadow(shape):
    shape.shadow.inherit = False


def rect(slide, x, y, w, h, color, alpha=100, name="Фигура", line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, U(x), U(y), U(max(w, 0.02)), U(max(h, 0.02)))
    no_shadow(sh)
    if line is None:
        sh.line.fill.background()
    else:
        lcolor, lalpha, lw = line
        sh.line.color.rgb = lcolor
        sh.line.width = Pt(lw)
        ln = sh._element.spPr.find(qn("a:ln"))
        srgb = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(srgb, qn("a:alpha")); a.set("val", str(int(lalpha * 1000)))
    if color is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = color
        if alpha < 100:
            srgb = sh.fill.fore_color._xFill.find(qn("a:srgbClr"))
            a = etree.SubElement(srgb, qn("a:alpha")); a.set("val", str(int(alpha * 1000)))
    sh.text_frame.text = ""
    name_shape(sh, name)
    return sh


def hairline(slide, x, y, w, h, hexcolor, alpha, name="Линия"):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, U(x), U(y), U(max(w, 0.02)), U(max(h, 0.02)))
    no_shadow(sh); sh.line.fill.background()
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(hexcolor)
    srgb = sh.fill.fore_color._xFill.find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha")); a.set("val", str(int(alpha * 1000)))
    sh.text_frame.text = ""
    name_shape(sh, name)
    return sh


def scrim_top(slide, name="ЗАТЕМНЕНИЕ (VIGNETTE)"):
    h = 11.0
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, U(0), U(0), U(100), U(h))
    no_shadow(sh); sh.line.fill.background()
    spPr = sh._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    xml = (
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">'
        '<a:gsLst>'
        '<a:gs pos="0"><a:srgbClr val="0A0908"><a:alpha val="44000"/></a:srgbClr></a:gs>'
        '<a:gs pos="100000"><a:srgbClr val="0A0908"><a:alpha val="0"/></a:srgbClr></a:gs>'
        '</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>')
    ln = spPr.find(qn("a:ln"))
    spPr.insert(list(spPr).index(ln) if ln is not None else len(spPr), etree.fromstring(xml))
    name_shape(sh, name)
    return sh


def picture(slide, x, y, w, h, path, name, descr=None, outline=None):
    pic = slide.shapes.add_picture(str(path), U(x), U(y), U(w), U(h))
    if outline:
        color, alpha, width_pt = outline
        pic.line.color.rgb = color
        pic.line.width = Pt(width_pt)
        ln = pic._element.spPr.find(qn("a:ln"))
        srgb = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(srgb, qn("a:alpha")); a.set("val", str(int(alpha * 1000)))
    name_shape(pic, name, descr)
    return pic


def text(slide, x, y, w, h, runs, size, family=FONT_SANS, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, ls_em=0.0, upper=False, bold=False, lh=None, name="Текст"):
    """runs: строка или список (текст, RGBColor)."""
    box = slide.shapes.add_textbox(U(x), U(y), U(w), U(max(h, 1)))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit"):
        el = bodyPr.find(qn(tag))
        if el is not None:
            bodyPr.remove(el)
    if isinstance(runs, str):
        runs = [(runs, C_SECONDARY)]
    lines = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
    spc = int(round(size * U_PT * ls_em * 100))
    for i, line_runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for txt, color in line_runs:
            if upper:
                txt = txt.upper()
            run = p.add_run(); run.text = txt
            run.font.size = PT(size); run.font.name = family; run.font.color.rgb = color
            run.font.bold = bold
            rPr = run._r.get_or_add_rPr()
            if spc:
                rPr.set("spc", str(spc))
            for tag in ("a:latin", "a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rPr, qn(tag))
                el.set("typeface", family)
        if lh:
            pPr = p._p.get_or_add_pPr()
            ln = etree.SubElement(pPr, qn("a:lnSpc"))
            pts = etree.SubElement(ln, qn("a:spcPts")); pts.set("val", str(int(size * U_PT * lh * 100)))
            pPr.insert(0, ln)
    name_shape(box, name)
    return box


def mark_and_folio(slide, fol_text, is_after):
    hairline(slide, 4.6, 5.05, 2.4, 0.05, "9A7F5E", 100, "Акцент-линия")
    text(slide, 7.4, 4.6, 30, 2, "Реставрация объектов", .58, FONT_SANS_SB,
         ls_em=.36, upper=True, name="Марка")
    text(slide, 65.4, 4.6, 30, 2, fol_text, .55, FONT_SANS_SB, align=PP_ALIGN.RIGHT,
         ls_em=.28, upper=True, name="Фолио")


def spec_grid(slide, top):
    labels = S.SPEC_LABELS
    n = len(labels)
    gap = 1.4
    col_w = (90.8 - gap * (n - 1)) / n
    hairline(slide, 4.6, top, 90.8, 0.05, "F2ECE3" if False else "26241D", 22, "Разделитель-специф")
    for i, lbl in enumerate(labels):
        x = 4.6 + i * (col_w + gap)
        text(slide, x, top + 1.1, col_w, 1.2, [[(lbl, C_TAUPE)]], .5, FONT_SANS, ls_em=.12, upper=True,
             name=f"Характеристика — {lbl}")
        hairline(slide, x, top + 2.55, col_w, 0.045, "26241D", 16, "Линия значения")


def strip(slide, o, is_after):
    name = o["name"].replace("\n", " ")
    if is_after:
        h = 15.6
    else:
        h = 9.2
    top = 56.25 - h
    bg = C_SECONDARY if is_after else C_PRIMARY
    rect(slide, 0, top, 100, h, bg, name="Полоса-подпись")
    row_top = top + (h - 4.6) / 2 if not is_after else top + 1.9
    text(slide, 4.6, row_top, 40, 1.3, [[(f"Объект {o['idx']}", C_ACCENT)]], .58, FONT_SANS_SB,
         ls_em=.36, upper=True, name="Киккер объекта")
    text(slide, 4.6, row_top + 1.6, 60, 3.2,
         [[(name, C_PRIMARY if is_after else C_SECONDARY)]], 2.35, FONT_SERIF, name="Название объекта")
    state = "После реставрации" if is_after else "До реставрации"
    state_color = C_ACCENT if is_after else C_STONE
    text(slide, 55.4, row_top + 0.35, 40, 1.3, [[(state, state_color)]], .55, FONT_SANS_SB,
         align=PP_ALIGN.RIGHT, ls_em=.28, upper=True, name="Статус")
    if is_after:
        spec_grid(slide, top + 9.1)


def hero_picture(slide, slot, label):
    if slot in PLACEHOLDER_SLOTS:
        path = ROOT / f"build/ph_e/{slot}.png"
        picture(slide, 0, 0, 100, 56.25, path, f"ФОТО-ПЛЕЙСХОЛДЕР — {slot}",
                f"Место под будущую фотографию: {label}. Заменить: правый клик — «Изменить рисунок».")
    else:
        path = ROOT / f"assets/e/{slot}.jpg"
        picture(slide, 0, 0, 100, 56.25, path, f"ФОТО — {slot}",
                f"Фотография: {label}. Заменить: правый клик — «Изменить рисунок».")


def object_before_slide(prs, o):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slot = f"{o['slug']}-before"
    hero_picture(slide, slot, f"{o['name'].replace(chr(10),' ')} — ДО")
    scrim_top(slide)
    mark_and_folio(slide, f"{o['idx']} / 06", False)
    strip(slide, o, False)
    return slide


def object_after_slide(prs, o):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slot = f"{o['slug']}-after"
    hero_picture(slide, slot, f"{o['name'].replace(chr(10),' ')} — ПОСЛЕ")
    scrim_top(slide)
    mark_and_folio(slide, f"{o['idx']} / 06", True)
    det_path = ROOT / f"assets/e/{slot}-det.jpg"
    if det_path.exists():
        ox, oy, ow, oh = 79.4, 8.5, 16.5, 22.0
        picture(slide, ox, oy, ow, oh, det_path, f"ДЕТАЛЬ ФАСАДА — {slot}",
                "Фрагмент фасада — деталь для крупного плана.",
                outline=(RGBColor.from_string("F2ECE3"), 40, 0.5))
        cap = "Портик после реставрации" if o["slug"] == "giliarovskogo" else "Фрагмент фасада"
        text(slide, 79.9, 31.2, 15.5, 1.2, cap, .55, FONT_SANS_SB, align=PP_ALIGN.RIGHT,
             ls_em=.28, upper=True, name="Подпись детали")
    strip(slide, o, True)
    return slide


def cover_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture(slide, 0, 0, 100, 56.25, ROOT / "assets/e/cover.jpg", "ФОТО — обложка",
            "Фотография обложки. Заменить: правый клик — «Изменить рисунок».")
    scrim_top(slide)
    text(slide, 4.6, 42, 70, 8,
         [[(S.TITLE, C_SECONDARY)]], 4.3, FONT_SERIF, ls_em=.16, upper=True,
         anchor=MSO_ANCHOR.BOTTOM, lh=1.05, name="Заголовок")
    text(slide, 4.6, 50.7, 70, 3,
         S.SUBTITLE, 1.15, FONT_SANS, name="Подзаголовок")
    return slide


def final_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture(slide, 0, 0, 100, 56.25, ROOT / "assets/e/final.jpg", "ФОТО — финал",
            "Итоговая фотография. Заменить: правый клик — «Изменить рисунок».")
    scrim_top(slide)
    mark_and_folio(slide, "Портфолио — 06 объектов", True)
    h = 7.4
    top = 56.25 - h
    rect(slide, 0, top, 100, h, C_SECONDARY, name="Полоса-индекс")
    n = len(S.OBJECTS)
    gap = 2.6
    pad = 1.3
    # приблизительная раскладка: считаем текстовые блоки равной шириной
    col_w = (90.8 - gap * (n - 1)) / n
    for i, o in enumerate(S.OBJECTS):
        x = 4.6 + i * (col_w + gap)
        if i > 0:
            hairline(slide, x - gap / 2, top + 1.6, 0.05, h - 3.2, "26241D", 22, "Разделитель")
        text(slide, x, top + 2.6, col_w, 2.2,
             [[(o["idx"] + "  ", C_ACCENT), (o["name"].split(chr(10))[0], C_TAUPE)]],
             .58, FONT_SANS_SB, ls_em=.05, name=f"Индекс — {o['slug']}")
    return slide


prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
cover_slide(prs)
for o in S.OBJECTS:
    object_before_slide(prs, o)
    object_after_slide(prs, o)
final_slide(prs)
prs.save(OUT)
print("OK", OUT, round(OUT.stat().st_size / 1024 / 1024, 2), "MB",
      "слайдов:", len(prs.slides._sldIdLst))
