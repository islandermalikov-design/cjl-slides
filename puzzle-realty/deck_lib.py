#!/usr/bin/env python3
"""Общие помощники для презентаций Puzzle Realty.

Собирается на шаблоне конгресса (assets/congress-template.potx): баннер сверху приходит
с мастер-слайда, поэтому весь контент лежит ниже TOP.
Использование: from deck_lib import *; start(); ...слайды...; finish(путь)
Шрифт Gilroy должен быть установлен на машине, где открывают файл.
"""
import os
import re
import shutil
import zipfile
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from PIL import Image
import qrcode

HERE = os.path.dirname(os.path.abspath(__file__))
PH = os.path.join(HERE, "assets", "photos")
LG = os.path.join(HERE, "assets", "logo")
TMP = os.path.join(HERE, "out", "tmp")
os.makedirs(TMP, exist_ok=True)

# ── фирменные токены ──────────────────────────────────────────────
FONT = "Gilroy"
RED = RGBColor(0xE7, 0x2F, 0x2A)
RED_DK = RGBColor(0xC4, 0x22, 0x1E)
PINK = RGBColor(0xFC, 0xE9, 0xE8)
PINK_2 = RGBColor(0xFA, 0xD5, 0xD4)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT = RGBColor(0xF4, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = 13.333, 7.5
TOP = 1.2  # ниже баннера конгресса из мастер-слайда


def _open_template():
    """python-pptx не открывает .potx — подменяем content-type на обычный .pptx."""
    src = os.path.join(HERE, "assets", "congress-template.potx")
    dst = os.path.join(TMP, "template.pptx")
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zo.writestr(it, data)
    p = Presentation(dst)
    for sid in list(p.slides._sldIdLst):  # пустой слайд-заготовка
        p.part.drop_rel(sid.rId)
        p.slides._sldIdLst.remove(sid)
    return p


prs = None
LAYOUT = None


def start():
    """Новая презентация на шаблоне конгресса."""
    global prs, LAYOUT
    prs = _open_template()
    LAYOUT = next(l for l in prs.slide_layouts if l.name == "Только заголовок")
    return prs


# ── помощники ─────────────────────────────────────────────────────
def _alpha(shape, pct):
    sf = shape.fill._xPr.find(qn("a:solidFill"))
    clr = sf[0]
    a = clr.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))})
    clr.append(a)


def nostyle(sh):
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         radius=None, alpha=None, dash=False):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if alpha is not None:
            _alpha(sh, alpha)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            ln = sh.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def rrect(s, x, y, w, h, fill=None, line=None, lw=1.0, r=0.12, alpha=None):
    # радиус задаётся в долях от меньшей стороны
    return rect(s, x, y, w, h, fill, line, lw, MSO_SHAPE.ROUNDED_RECTANGLE,
                radius=min(0.5, r / min(w, h)), alpha=alpha)


def text(s, x, y, w, h, paras, size=16, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.05, after=0, caps=False, italic=False):
    """paras: str | list[str | dict(text,size,color,bold,after,align,italic,runs)]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        if isinstance(p, str):
            p = {"text": p}
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        para.line_spacing = p.get("spacing", spacing)
        para.space_after = Pt(p.get("after", after))
        runs = p.get("runs") or [dict(text=p["text"])]
        for r in runs:
            run = para.add_run()
            t = r["text"]
            run.text = t.upper() if p.get("caps", caps) else t
            f = run.font
            f.name = FONT
            f.size = Pt(r.get("size", p.get("size", size)))
            f.bold = r.get("bold", p.get("bold", bold))
            f.italic = r.get("italic", p.get("italic", italic))
            f.color.rgb = r.get("color", p.get("color", color))
    return tb


def pic(s, path, x, y, w, h, focus=(0.5, 0.5)):
    """Картинка «cover»: заполняет рамку, лишнее обрезается вокруг точки focus."""
    iw, ih = Image.open(path).size
    box, img = w / h, iw / ih
    p = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if img > box:  # шире рамки — режем по бокам
        vis = box / img
        left = min(max(focus[0] - vis / 2, 0), 1 - vis)
        p.crop_left, p.crop_right = left, 1 - vis - left
    else:
        vis = img / box
        top = min(max(focus[1] - vis / 2, 0), 1 - vis)
        p.crop_top, p.crop_bottom = top, 1 - vis - top
    return p


def logo(s, x, y, w, white=False):
    path = os.path.join(LG, "logo_white.png" if white else "logo_dark.png")
    iw, ih = Image.open(path).size
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(w * ih / iw))


def title(s, t, x=0.9, y=1.5, w=11.5, size=30, color=DARK, bar=RED, kicker=None, lines=None):
    """Заголовок слайда — настоящий title-плейсхолдер + фирменная красная черта."""
    lines = lines or 1 + int(len(t) * size * 0.0092 / w)
    hh = lines * size * 1.2 / 72 + 0.05
    if kicker:
        text(s, x, y - 0.32, w, 0.28, kicker, size=12, color=RED, bold=True, caps=True)
    if bar is not None:
        rect(s, x - 0.28, y - 0.03, 0.07, hh + 0.06, fill=bar)
    ph = s.shapes.title
    el = ph._element
    el.getparent().remove(el)
    s.shapes._spTree.append(el)  # поверх фоновых плашек
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(hh)
    tf = ph.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.line_spacing = 1.0
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = t.upper()
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
    r.font.color.rgb = color
    return hh


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def slide():
    # фон слайда закрыт картинкой мастера, поэтому цветные поля рисуем плашками
    return prs.slides.add_slide(LAYOUT)


def bullets(s, x, y, w, items, size=15, color=DARK, gap=7, dot=RED):
    """Список с красными ромбами-маркерами; возвращает высоту."""
    cy = y
    for it in items:
        lines = 1 + int(len(it) * size * 0.0078 / (w - 0.35))
        hh = lines * size * 1.25 / 72
        rect(s, x, cy + 0.07, 0.09, 0.09, fill=dot, shape=MSO_SHAPE.DIAMOND)
        text(s, x + 0.28, cy, w - 0.28, hh, it, size=size, color=color, spacing=1.05)
        cy += hh + gap / 72
    return cy - y


def pill(s, x, y, w, h, t, size=13, fill=RED, color=WHITE):
    rrect(s, x, y, w, h, fill=fill, r=h / 2)
    text(s, x, y, w, h, t, size=size, color=color, bold=True, caps=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def crop_file(src, box, name):
    im = Image.open(os.path.join(PH, src))
    iw, ih = im.size
    out = os.path.join(TMP, name)
    im.crop((int(box[0] * iw), int(box[1] * ih), int(box[2] * iw), int(box[3] * ih))).save(out, quality=92)
    return out




def tg_qr(path=None):
    """QR на Telegram-канал @puzzlerealty (из контактов прежней презентации)."""
    path = path or os.path.join(TMP, "qr_tg.png")
    q = qrcode.QRCode(border=1, box_size=10)
    q.add_data("https://t.me/puzzlerealty"); q.make(fit=True)
    q.make_image(fill_color="black", back_color="white").save(path)
    return path


def finish(out_path):
    """Сохраняет файл и прописывает в тему шрифт Gilroy и фирменные цвета."""
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    raw = os.path.join(TMP, "raw.pptx")
    prs.save(raw)
    BRAND = {"accent1": "E72F2A", "accent2": "1B3A8C", "accent3": "E6C88F", "accent4": "FAD5D4",
             "accent5": "C4221E", "accent6": "6B6B6B", "dk2": "1A1A1A", "lt2": "FCE9E8"}
    with zipfile.ZipFile(raw) as zi, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("ppt/theme/theme") and it.filename.endswith(".xml"):
                t = data.decode("utf8")
                t = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*', r"\g<1>" + FONT, t)
                for k, v in BRAND.items():
                    t = re.sub(rf'(<a:{k}>)<a:srgbClr val="[0-9A-Fa-f]{{6}}"', rf'\g<1><a:srgbClr val="{v}"', t)
                data = t.encode("utf8")
            zo.writestr(it, data)
    print("saved", out_path)
