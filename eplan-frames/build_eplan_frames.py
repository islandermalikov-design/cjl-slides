# Переносит рамки EPLAN (А3, ГОСТ 2.104) в шаблон PowerPoint.
# Координаты рамки взяты из векторов PDF EPLAN (в пунктах, 1 pt = 12700 EMU).
import copy, sys
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, DST = sys.argv[1], sys.argv[2]
P = lambda v: Emu(int(round(v * 12700)))
W_PT, H_PT = 1190.55, 841.89          # А3 альбомная, 420 × 297 мм
LW = 0.99                              # толщина линий EPLAN, pt
FONT = 'Tahoma'
BLACK = RGBColor(0, 0, 0)

prs = Presentation(SRC)
OLD_W, OLD_H = prs.slide_width, prs.slide_height
prs.slide_width, prs.slide_height = P(W_PT), P(H_PT)
prs.part._element.find(qn('p:sldSz')).attrib.pop('type', None)


# ---------- примитивы ----------
def line(g, x1, y1, x2, y2, dashed=False):
    c = g.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, P(x1), P(y1), P(x2), P(y2))
    c.line.width = Pt(LW)
    c.line.color.rgb = BLACK
    ln = c.line._get_or_add_ln()
    st = c._element.find(qn('p:style'))  # без стиля темы: иначе PowerPoint добавляет тень
    if st is not None:
        c._element.remove(st)
    if dashed:  # штрих 5.67 pt / пробел 5.67 pt, как в EPLAN
        cd = etree.SubElement(ln, qn('a:custDash'))
        ds = etree.SubElement(cd, qn('a:ds'))
        ds.set('d', '573000'); ds.set('sp', '573000')
    return c


def text(g, x0, y0, x1, y1, s, size, align='c', rot=False, bold=False,
         anchor=MSO_ANCHOR.MIDDLE, pad=0.0, name=None):
    """Текст в ячейке (x0,y0)-(x1,y1); rot=True — снизу вверх, как в боковой графе."""
    w, h = x1 - x0, y1 - y0
    if rot:
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        w, h = h, w
        x0, y0 = cx - w / 2, cy - h / 2
    tb = g.shapes.add_textbox(P(x0), P(y0), P(w), P(h))
    if name:
        tb.name = name
    if rot:
        tb.rotation = 270
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    tf.margin_left = tf.margin_right = P(pad)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = {'c': PP_ALIGN.CENTER, 'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT}[align]
    r = p.add_run()
    r.text = s
    f = r.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = BLACK
    rpr = r._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        e = etree.SubElement(rpr, qn(tag)); e.set('typeface', FONT)
    return tb


def new_group(slide, name):
    g = slide.shapes.add_group_shape()
    g.name = name
    return g


def to_back(slide, shape):
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(2, shape._element)


# ---------- элементы рамки ----------
def outer_frame(g):
    # внутренняя рамка поля чертежа
    line(g, 56.7, 14.3, 1176.4, 14.3)
    line(g, 56.7, 827.8, 1176.4, 827.8)
    line(g, 56.7, 14.3, 56.7, 827.8)
    line(g, 1176.4, 14.3, 1176.4, 827.8)
    text(g, 891.5, 830.0, 935.3, 841.8, 'Копировал', 8.8)
    text(g, 1060.0, 830.0, 1146.0, 841.8, 'Формат  А3', 8.8)


def left_strip_bottom(g, labels, size):
    """Нижняя боковая графа: 5 ячеек (Подп. и дата … Инв. № подл)."""
    ys = [416.8, 516.0, 586.9, 657.7, 757.0, 827.8]
    line(g, 22.7, 416.8, 22.7, 827.8)
    line(g, 36.9, 416.8, 36.9, 827.8)
    for y in ys:
        line(g, 22.7, y, 56.7, y)
    for (y0, y1), s in zip(zip(ys, ys[1:]), labels):
        text(g, 22.7, y0, 36.9, y1, s, size, rot=True)


def left_strip_top(g):
    """Верхняя боковая графа (штриховая): Перв. примен / Справ. №."""
    line(g, 22.7, 14.3, 22.7, 354.4, dashed=True)
    line(g, 36.9, 14.3, 36.9, 354.4, dashed=True)
    for y in (14.3, 184.4, 354.4):
        line(g, 22.7, y, 56.7, y, dashed=True)
    text(g, 22.7, 14.3, 36.9, 184.4, 'Перв. примен', 8.8, rot=True)
    text(g, 22.7, 184.4, 36.9, 354.4, 'Справ. №', 8.8, rot=True)


def column_header(g):
    """Строка номеров колонок 0…9 над полем чертежа."""
    xs = [56.7 + i * (1176.4 - 56.7) / 10 for i in range(11)]
    for x in xs[1:-1]:
        line(g, x, 0.5, x, 14.3)
    for i, (x0, x1) in enumerate(zip(xs, xs[1:])):
        text(g, x0, 3.5, x1, 14.3, str(i), 8.8)


def page_border(g):
    h = LW / 2  # линия по краю листа, утопленная на полтолщины, чтобы не обрезалась
    line(g, h, h, W_PT - h, h)
    line(g, h, H_PT - h, W_PT - h, H_PT - h)
    line(g, h, h, h, H_PT - h)
    line(g, W_PT - h, h, W_PT - h, H_PT - h)


def stamp_short(g, sheet, designation='00000'):
    """Основная надпись для последующих листов (форма 2а)."""
    X0, X1, XN, Y0, Y1 = 652.0, 1176.4, 836.2, 785.3, 827.8
    line(g, X0, Y0, X1, Y0)
    line(g, X0, Y0, X0, Y1)
    for y in (799.5, 813.7):
        line(g, X0, y, XN, y)
    cols = [652.0, 671.8, 700.2, 765.4, 807.9, 836.2]
    for x in cols[1:]:
        line(g, x, Y0, x, Y1)
    line(g, 1148.0, Y0, 1148.0, Y1)
    line(g, 1148.0, 805.1, X1, 805.1)
    for (a, b), s in zip(zip(cols, cols[1:]), ['Ред.', 'Листов', '№ докум', 'Подп', 'Дата']):
        text(g, a, 813.7, b, 827.8, s, 7.04)
    text(g, 1148.0, Y0, X1, 805.1, 'Лист', 8.8)
    text(g, 1148.0, 805.1, X1, Y1, str(sheet), 8.8, name='Номер листа')
    text(g, XN, Y0, 1148.0, Y1, designation, 16, name='Обозначение')


def stamp_big(g, logo_blob, sheet_name, designation='00000', stage='Р',
              sheet='—', sheets='—', title=''):
    """Основная надпись для первого листа (форма 1)."""
    X0, XM, X1, Y0, Y1 = 652.0, 836.2, 1176.4, 671.9, 827.8
    line(g, X0, Y0, X1, Y0)
    line(g, X0, Y0, X0, Y1)
    line(g, X1, Y0, X1, Y1)
    for y in (686.1, 714.4, 728.6, 757.0, 771.1, 785.3, 799.5, 813.7):
        line(g, X0, y, XM, y)
    for y in (700.3, 742.8):
        line(g, X0, y, X1, y)
    for x in (671.8, 737.0):
        line(g, x, Y0, x, 742.8)
    for x in (708.7, 765.4, 807.9, XM):
        line(g, x, Y0, x, Y1)
    line(g, 1034.7, 742.8, 1034.7, Y1)
    line(g, 1034.7, 757.0, X1, 757.0)
    line(g, 1034.7, 785.3, X1, 785.3)
    for x in (1077.2, 1125.3):
        line(g, x, 742.8, x, 785.3)
    cols = [652.0, 671.8, 708.7, 737.0, 765.4, 807.9, 836.2]
    for (a, b), s in zip(zip(cols, cols[1:]), ['Изм.', 'Кол.уч.', 'Лист', '№док.', 'Подп.', 'Дата']):
        text(g, a, 728.6, b, 742.8, s, 7.04)
    rows = [742.8, 757.0, 771.1, 785.3, 799.5, 813.7, 827.8]
    for (a, b), s in zip(zip(rows, rows[1:]), ['Разработал', 'Проверил', '', '', 'Н.контроль', 'Утвердил']):
        if s:
            text(g, 652.0, a, 708.7, b, s, 7.04, align='l', pad=2.8)
    hx = [1034.7, 1077.2, 1125.3, 1176.4]
    for (a, b), s in zip(zip(hx, hx[1:]), ['Стадия', 'Лист', 'Листов']):
        text(g, a, 742.8, b, 757.0, s, 8.8)
    for (a, b), s, nm in zip(zip(hx, hx[1:]), [stage, sheet, sheets], ['Стадия', 'Лист', 'Листов']):
        text(g, a, 757.0, b, 785.3, s, 8.8, name='Поле ' + nm)
    text(g, XM, Y0, X1, 700.3, designation, 12.32, name='Обозначение')
    text(g, XM, 700.3, X1, 742.8, title, 12.32, name='Наименование изделия')
    text(g, XM, 785.3, 1034.7, Y1, sheet_name, 8.8, name='Наименование документа')
    # логотип и реквизиты
    pic = g.shapes.add_picture(logo_blob, P(1054.5), P(788.1), P(99.2), P(23.1))
    pic.name = 'Логотип'
    text(g, 1080.0, 811.0, 1173.0, 817.0, 'IntelMet Technologies Ltd.', 4.75, align='r')
    text(g, 1080.0, 818.3, 1173.0, 824.3, '199178, Санкт-Петербург', 4.75, align='r')


STRIP_SHORT = ['Подп. и дата', 'Инв. № дубл', 'Взам. инв. №', 'Подп. и дата', 'Инв. № подл']
STRIP_BIG = ['Подпись и дата', 'Инв.№  дубл.', 'Взамен инв.№', 'Подпись и дата', 'Инв.№  подп.']


def frame_sheet(slide, sheet):
    g = new_group(slide, 'Рамка EPLAN (А3, лист)')
    page_border(g)
    column_header(g)
    outer_frame(g)
    left_strip_top(g)
    left_strip_bottom(g, STRIP_SHORT, 8.8)
    stamp_short(g, sheet)
    to_back(slide, g)


def frame_first(slide, logo, sheet_name):
    g = new_group(slide, 'Рамка EPLAN (А3, первый лист)')
    outer_frame(g)
    left_strip_bottom(g, STRIP_BIG, 7.04)
    stamp_big(g, logo, sheet_name)
    to_back(slide, g)


def frame_title(slide):
    g = new_group(slide, 'Рамка EPLAN (А3, титул)')
    outer_frame(g)
    left_strip_bottom(g, STRIP_BIG, 7.04)
    to_back(slide, g)


# ---------- перенос содержимого ----------
def move_content(slide, ox, oy, nx, ny, s):
    """Масштабирует и сдвигает содержимое: (ox,oy) старого слайда → (nx,ny) нового, дюймы."""
    for sh in slide.shapes:
        sh.left = int(Inches(nx) + (sh.left - Inches(ox)) * s)
        sh.top = int(Inches(ny) + (sh.top - Inches(oy)) * s)
        sh.width = int(sh.width * s)
        sh.height = int(sh.height * s)
        el = sh._element
        for e in el.iter():
            if e.tag in (qn('a:rPr'), qn('a:defRPr'), qn('a:endParaRPr')) and e.get('sz'):
                e.set('sz', str(int(round(int(e.get('sz')) * s / 50.0)) * 50))
            elif e.tag == qn('a:gridCol'):
                e.set('w', str(int(int(e.get('w')) * s)))
            elif e.tag == qn('a:tr'):
                e.set('h', str(int(int(e.get('h')) * s)))
            elif e.tag in (qn('a:bodyPr'), qn('a:tcPr')):
                for k in ('lIns', 'rIns', 'tIns', 'bIns', 'marL', 'marR', 'marT', 'marB'):
                    if e.get(k):
                        e.set(k, str(int(int(e.get(k)) * s)))


def remove(slide, names):
    for sh in list(slide.shapes):
        if sh.name in names:
            sh._element.getparent().remove(sh._element)


def heading(slide, s):
    tb = slide.shapes.add_textbox(P(88.0), P(18.0), P(560.0), P(20.0))
    tb.name = 'Наименование листа'
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    r = tf.paragraphs[0].add_run()
    r.text = s
    r.font.name, r.font.size, r.font.bold, r.font.underline = FONT, Pt(16), True, True
    r.font.color.rgb = BLACK


S = 1.2
slides = list(prs.slides)

# 1 — титульный лист
s1 = slides[0]
logo_blob = None
for sh in s1.shapes:
    if sh.shape_type == 13:
        import io
        logo_blob = io.BytesIO(sh.image.blob)
remove(s1, {'Rectangle 1'})
move_content(s1, 0.65, 0.48, 1.45, 1.6, S)
frame_title(s1)

# 2 — ведомость: первый лист, основная надпись формы 1
s2 = slides[1]
move_content(s2, 0.4, 0.22, 1.05, 0.45, S)
logo_blob.seek(0)
frame_first(s2, logo_blob, 'Ведомость основных комплектов')

# 3…5 — листы схем: основная надпись формы 2а, строка колонок 0…9
old_frame = {'Rectangle 1', 'Rectangle 2', 'Rectangle 13', 'Rectangle 14', 'Rectangle 18',
             'Rectangle 21', 'Rectangle 24'} | {f'TextBox {i}' for i in range(3, 29)}
for n, (sl, name) in enumerate(zip(slides[2:], ['Наименование листа', 'Условные обозначения',
                                                 'Структурная схема']), start=1):
    remove(sl, old_frame)
    move_content(sl, 0.32, 0.28, 1.05, 0.75, S)
    heading(sl, name)
    frame_sheet(sl, n)

prs.save(DST)
print('saved', DST)
